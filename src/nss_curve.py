import numpy as np
from scipy.optimize import least_squares
from matplotlib import pyplot as plt
from itertools import product

import pandas as pd
from pathlib import Path
from openpyxl import load_workbook

# BEFORE RUNNING THE SCRIPT
# This script can be run in two different ways:
#   1) Grid search for optimal parameters (slow)
#   2) Run with good initial guess (fast) (useful
#      to get the plots and final results)
# The first choice should be just used once, and
# its results copied so the second one can work.


# GET THE PATH TO THE DATA FILES
rootPath = Path(__file__).resolve().parent
readingFileName = "AIRMM_V2.xlsm"
printingFileName = "AIRMM_V2.xlsm"
readingDataPath = rootPath / readingFileName
printingDataPath = rootPath / printingFileName

# Calculate NSS, given the parameters and timestamps
def NSS_vectorial(prms, times):
    b0, b1, b2, b3, t1, t2 = prms
    tt1 = times/t1; tt2 = times/t2
    term1 = b0 + b1 * (1-np.exp(-tt1))/tt1
    term2 = b2 * ((1-np.exp(-tt1))/tt1 - np.exp(-tt1))
    term3 = b3 * ((1-np.exp(-tt2))/tt2 - np.exp(-tt2))
    return term1 + term2 + term3

# Objective function for least squares minimization
def objective(prms, times, yields):
    return NSS_vectorial(prms, times) - yields

# Since we are trying to minimize the MSE numerically, we need a starting point
# x0 for the numerical method to work. The NSS curve is very sensitive to the
# initial data; thus, we will run the algorithm on several starting points to
# get the optimal result.
def grid_search_EURIBOR(times, yields):
    # Total of guesses = 2 * 2 * 2 * 2 * 2 * 2 = 64
    # We take an educated guess on the inital parameters based on the plotted
    # curve of df["actualTime"] vs. df["actualRates"]
    b0_grid = np.linspace(200, 250, 2)
    b1_grid = np.linspace(-50, 50, 2)
    b2_grid = np.linspace(-50, 50, 2)
    b3_grid = np.linspace(-50, 50, 2)
    t1_grid  = np.linspace(0.1, 15, 2)
    t2_grid  = np.linspace(10, 50, 2)
    optimal_prms = np.array([0, 0, 0, 0, 0, 0])
    best_fit = np.inf
    count = 0
    count_best_fit = 0
    for b0, b1, b2, b3, t1, t2 in product(b0_grid, b1_grid, b2_grid, b3_grid,
                                            t1_grid, t2_grid):
        initial_prms = [b0, b1, b2, b3, t1, t2]
        reg = least_squares(objective, initial_prms, 
                args=(times, yields))
        fit = reg.cost
        count +=1
        print(f"Count: {count}")
        # We check that the fit we just found, complies with the restrictions
        # of the NSS model
        if fit<best_fit and reg.x[4]>0 and reg.x[5]>0:
            best_fit = fit
            optimal_prms = reg.x
            print(f"New optimal parametres: {count_best_fit}")
            count_best_fit += 1
    return optimal_prms

def grid_search_OIS(times, yields):
    # Total of guesses = 2 * 3 * 2 * 2 * 2 * 2 = 96
    # We take an educated guess on the inital parameters based on the plotted
    # curve of df["actualTime"] vs. df["actualRates"]
    b0_grid = np.linspace(180, 230, 2)
    b1_grid = np.linspace(-100, 100, 3)
    b2_grid = np.linspace(-50, 50, 2)
    b3_grid = np.linspace(-50, 50, 2)
    t1_grid  = np.linspace(0.1, 15, 2)
    t2_grid  = np.linspace(10, 50, 2)
    optimal_prms = np.array([0, 0, 0, 0, 0, 0])
    best_fit = np.inf
    count = 0
    count_best_fit = 0
    for b0, b1, b2, b3, t1, t2 in product(b0_grid, b1_grid, b2_grid, b3_grid,
                                            t1_grid, t2_grid):
        initial_prms = [b0, b1, b2, b3, t1, t2]
        reg = least_squares(objective, initial_prms, 
                args=(times, yields))
        fit = reg.cost
        count +=1
        print(f"Count: {count}")
        # We check that the fit we just found, complies with the restrictions
        # of the NSS model
        if fit<best_fit and reg.x[4]>0 and reg.x[5]>0:
            best_fit = fit
            optimal_prms = reg.x
            print(f"New optimal parametres: {count_best_fit}")
            count_best_fit += 1
    return optimal_prms

def print_curve_OIS(prms, dates, yields):
    timestamps = np.linspace(0.01, 60, int(1e4))
    fitted_curve = NSS_vectorial(prms, timestamps)
    plt.scatter(timestamps, fitted_curve, marker='o', 
                s=1, label = "Adjusted curve")
    plt.scatter(dates, yields, marker='x', s=15, 
                color='red', label="Historical data")
    plt.xlabel('Time (years)')
    plt.ylabel('EUR OIS')
    plt.title("NSS model curve adjusted to current EUR OIS")
    plt.legend()
    plt.grid(True)
    plt.show()
    
def print_curve_EURIBOR(prms, dates, yields):
    timestamps = np.linspace(0.01, 60, int(1e4))
    fitted_curve = NSS_vectorial(prms, timestamps)
    plt.scatter(timestamps, fitted_curve, marker='o', 
                s=1, label = "Adjusted curve")
    plt.scatter(dates, yields, marker='x', s=15, 
                color='red', label="Historical data")
    plt.xlabel('Time (years)')
    plt.ylabel('EURIBOR6M')
    plt.title("NSS model curve adjusted to current EURIBOR6M")
    plt.legend()
    plt.grid(True)
    plt.show()

if __name__=="__main__":
    # ····· READING THE DATASET ·····
    df = pd.read_excel(
        readingDataPath,
        sheet_name="Yield Curves",
        header=1,
        usecols=[4, 7, 17, 20])
    df.columns = ["timeToMaturityOIS", "discountRateOIS",
                "timeToMaturityEURIBOR", "discountRateEURIBOR"]
    df["actualTimeOIS"] = df["timeToMaturityOIS"].dropna() / 365
    df["actualRatesOIS"] = df["discountRateOIS"].dropna()*1e4
    df["actualTimeEURIBOR"] = df["timeToMaturityEURIBOR"].dropna() / 365
    df["actualRatesEURIBOR"] = df["discountRateEURIBOR"].dropna()*1e4
    
    # OPTION 1: Search for the optimal parameters from scratch. Recommended
    #           just for the first time, to calibrate the curves. Otherwise,
    #           it's a very slow process
    #
    # optimal_prms_OIS = grid_search_OIS(df["actualTimeOIS"],
    #                                   df["actualRatesOIS"])
    # optimal_prms_EURIBOR = grid_search_EURIBOR(df["actualTimeEURIBOR"],
    #                                           df["actualRatesEURIBOR"])
    #
    # OPTION 2: Once the optimal parameters have been found, you can copy-paste
    #           them inside the vector 'initial_prms_x' and run just one 
    #           regression, instead of a full grid search. This is useful 
    #           in case we just want the plots or the parameters, once the
    #           initial big search has been executed
    #
    # OIS REGRESSION
    initial_prms_OIS = [195.94896210444642-200, -27.51791648811566,
                    23072.465626943027, -22971.986987337776,
                    6.574509621372123, 6.513383121629629]
    regression_OIS = least_squares(objective, initial_prms_OIS, 
                            args=(df["actualTimeOIS"].dropna(),
                                df["actualRatesOIS"].dropna()))
    optimal_prms_OIS = regression_OIS.x
    print(f"Optimal parameters OIS:\n
        b0={optimal_prms_OIS[0]}\nb1={optimal_prms_OIS[1]}\
        \nb2={optimal_prms_OIS[2]}\nb3={optimal_prms_OIS[3]}\
        \nt1={optimal_prms_OIS[4]}\nt2={optimal_prms_OIS[5]}\n")
    print_curve_OIS(optimal_prms_OIS, df["actualTimeOIS"], df["actualRatesOIS"])
    
    # EURIBOR REGRESSION
    initial_prms_EURIBOR = [195.94896210444642-200, -27.51791648811566,
                    23072.465626943027, -22971.986987337776,
                    6.574509621372123, 6.513383121629629]
    regression_EURIBOR = least_squares(objective, initial_prms_EURIBOR, 
                args=(df["actualTimeEURIBOR"], df["actualRatesEURIBOR"]))
    optimal_prms_EURIBOR = regression_EURIBOR.x
    print(f"Optimal parameters EURIBOR:\n
        b0={optimal_prms_EURIBOR[0]}\nb1={optimal_prms_EURIBOR[1]}\
        \nb2={optimal_prms_EURIBOR[2]}\nb3={optimal_prms_EURIBOR[3]}\
        \nt1={optimal_prms_EURIBOR[4]}\nt2={optimal_prms_EURIBOR[5]}\n")
    print_curve_EURIBOR(optimal_prms_EURIBOR, 
                        df["actualTimeEURIBOR"], df["actualRatesEURIBOR"])
    ##########################################################################
    
    # Write the computed rates into the excel again
    # 1) Load values (computed results) without formulas
    wb_vals = load_workbook(printingDataPath, keep_vba=False, data_only=True)
    ws_vals = wb_vals["Interpolated Rates"]
    scheduleSpot = [ws_vals[f"F{row}"].value/365 for row in range(5, 16)]
    scheduleSpot2 = [ws_vals[f"M{row}"].value/365 for row in range(5, 16)]

    # CALCULATE THE DESIRED RATES
    scheduleOisSpot = NSS_vectorial(optimal_prms_OIS, scheduleSpot)
    scheduleOisSpot2 = NSS_vectorial(optimal_prms_OIS, scheduleSpot2)
    scheduleEURIBORSpot = NSS_vectorial(optimal_prms_EURIBOR, scheduleSpot)
    scheduleEURIBORSpot2 = NSS_vectorial(optimal_prms_EURIBOR, scheduleSpot2)
    
    # 2) Reload file preserving VBA (and formulas)
    wb = load_workbook(printingDataPath, keep_vba=True, data_only=False)
    ws = wb["Interpolated Rates"]
    for row in range(5, 16):  # 5..15 included
        ws[f"H{row}"].value = float(scheduleOisSpot[row-5])
        ws[f"I{row}"].value = float(scheduleEURIBORSpot[row-5])
        ws[f"O{row}"].value = float(scheduleOisSpot2[row-5])
        ws[f"P{row}"].value = float(scheduleEURIBORSpot2[row-5])
    wb.save(printingDataPath)
